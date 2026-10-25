"""Word specification. Section order: docs/06-export.md#word-specification."""

from __future__ import annotations

from io import BytesIO

from docx import Document

from .. import cognition, vehicle
from ..knowledge import load_pack
from ..models import Requirement
from . import Handoff

TITLE = "In-Car Voice Assistant with Cognitive Intelligence — System Requirements Specification (SEW1)"


def _requirement(doc: Document, r: Requirement) -> None:
    doc.add_heading(f"{r.id} — {r.title}", level=3)
    doc.add_paragraph(r.statement)
    meta = [f"State: {r.state}", f"Priority: {r.priority}", f"Source: {r.source}"]
    if r.cognitive_capability:
        meta.append(f"Cognitive capability: {', '.join(r.cognitive_capability)}")
    if r.preliminary_asil:
        meta.append(f"Preliminary ASIL: {r.preliminary_asil} ({r.asil_note})")
    if r.standard_references:
        meta.append(f"References: {', '.join(r.standard_references)}")
    if r.scenario:
        parts = vehicle.render_parts(r.scenario)
        for kw in ("given", "when", "then"):
            para = doc.add_paragraph()
            para.add_run(f"{kw.upper()} ").bold = True
            para.add_run(parts[kw])
    doc.add_paragraph(f"Rationale: {r.rationale}")
    doc.add_paragraph(" · ".join(meta))


def _group(doc: Document, heading: str, reqs: list[Requirement], note: str = "") -> None:
    doc.add_heading(heading, level=1)
    if note:
        doc.add_paragraph(note)
    for r in reqs:
        _requirement(doc, r)
    if not reqs:
        doc.add_paragraph("No included requirements.")


def _table(doc: Document, header: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, header):
        cell.text = text
    for row in rows:
        for cell, text in zip(table.add_row().cells, row):
            cell.text = str(text)


def build(h: Handoff) -> bytes:
    by = lambda *cats: [r for r in h.requirements if r.category in cats]  # noqa: E731
    doc = Document()
    doc.add_heading(TITLE, level=0)

    doc.add_heading("1. Document control", level=1)
    counts = ", ".join(f"{k}: {v}" for k, v in sorted(h.state_counts.items())) or "none"
    for line in (f"Program: {h.project['name']}", f"Exported: {h.exported_at.isoformat(timespec='seconds')}",
                 "Authoring tool: SEW1 workbench", f"State counts: {counts}",
                 f"Cognitive level: target {h.target_level}, specified {h.specified_level}"):
        doc.add_paragraph(line)

    doc.add_heading("2. Scope", level=1)
    doc.add_paragraph("An in-car voice assistant specified as an AI partner and companion. Cognitive levels:")
    for level in load_pack("companion_cognition")["levels"]:
        doc.add_paragraph(f"{level['id']} {level['name']}", style="List Bullet")
    doc.add_paragraph("Companion principles:")
    for p in load_pack("companion_cognition")["principles"]:
        doc.add_paragraph(p, style="List Bullet")

    doc.add_heading("3. Context assumptions", level=1)
    doc.add_paragraph("A speech capture and playback path exists; vehicle signals reach the assistant through a defined "
                      "interface; a display is subject to driver-distraction rules. Deployment: "
                      f"{h.project.get('deployment_assumption', 'open')} (open assumption, not a design choice).")

    doc.add_heading("4. Vehicle context", level=1)
    cat = vehicle.catalog()
    doc.add_paragraph(f"Signal catalog: {cat['source']} ({cat['version']}), {cat['status']}. Signals use COVESA VSS "
                      "paths with an indicative AAOS VHAL mapping; the runtime choice stays open.")
    names = {d["id"]: d["name"] for d in cat["domains"]}
    rows = []
    for d in cat["domains"]:
        sigs = [s for s in h.signals if s["domain"] == d["id"]]
        rows.append([names[d["id"]], str(len(sigs)), str(sum(1 for s in sigs if s["access"] == "actuator")),
                     str(sum(1 for s in sigs if s["used_by"]))])
    _table(doc, ["Domain", "Signals", "Actuators", "Used by requirements"], rows)
    doc.add_paragraph("Triggers:")
    for t in cat["triggers"]:
        doc.add_paragraph(f"{t['id']} — {t['label']}", style="List Bullet")
    doc.add_paragraph("Responses:")
    for r in cat["responses"]:
        doc.add_paragraph(f"{r['type']} — {r['rule']}", style="List Bullet")
    doc.add_paragraph(f"{len(h.scenarios)} behavioural requirements carry a GIVEN / WHEN / THEN scenario, printed "
                      "with each requirement below and listed in the Scenarios sheet of the workbook.")

    _group(doc, "5. Stakeholder needs", by("STK"))
    _group(doc, "6. System requirements by category", by("SYS-FUNC", "SYS-PERF", "SYS-HMI"))

    doc.add_heading("7. Cognitive companion", level=1)
    cog = by("SYS-COG")
    level_of = {c["id"]: c["level"] for c in cognition.catalog()}
    for level in cognition.LEVELS[1:]:
        for r in [r for r in cog if min((level_of.get(c, "CL5") for c in r.cognitive_capability), default="CL5") == level]:
            _requirement(doc, r)
    _table(doc, ["Capability", "Level", "Covered", "Requirements"],
           [[f"{c['id']} {c['title']}", c["level"], "yes" if c["covered"] else "no", ", ".join(c["requirement_ids"])]
            for c in h.coverage])

    _group(doc, "8. Privacy", by("SYS-PRIV"), "Lawful basis remains subject to privacy review.")
    _group(doc, "9. Driver distraction", by("SYS-DIST"))
    _group(doc, "10. Safety and cybersecurity", by("SYS-SAFE", "SYS-SEC"),
           "Preliminary ASIL values are to be confirmed by the safety manager.")

    doc.add_heading("11. Verification matrix", level=1)
    _table(doc, ["ID", "Method", "Acceptance criteria", "State"],
           [[r.id, r.verification_method, r.acceptance_criteria, str(r.state)] for r in h.requirements])

    doc.add_heading("12. Open points", level=1)
    _table(doc, ["Kind", "ID", "Summary", "Owner"], [[o["kind"], o["id"], o["summary"], o["owner"]] for o in h.open_items])

    doc.add_heading("13. Change record", level=1)
    changed = [r for r in h.requirements if r.version > 1]
    for r in changed:
        doc.add_paragraph(f"{r.id} v{r.version}: {r.change_note}", style="List Bullet")
    if not changed:
        doc.add_paragraph("No versions greater than 1.")

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def filename(h: Handoff) -> str:
    return f"SEW1-IVA-specification-{h.stamp()}.docx"
