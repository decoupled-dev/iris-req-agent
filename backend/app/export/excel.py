"""Excel workbook. Sheet layout: docs/06-export.md#excel-workbook."""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

from . import Handoff

REQ_COLUMNS = ["ID", "Version", "State", "Title", "Statement", "Rationale", "Source", "Category",
               "Cognitive capability", "Priority", "Preliminary ASIL", "ASIL note", "Standard references",
               "Verification method", "Acceptance criteria", "Change note", "Updated"]


def _sheet(wb: Workbook, title: str, header: list[str], rows: list[list], first: bool = False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(header)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    return ws


def build(h: Handoff) -> bytes:
    wb = Workbook()
    _sheet(wb, "Requirements", REQ_COLUMNS, [
        [r.id, r.version, str(r.state), r.title, r.statement, r.rationale, r.source, r.category,
         ", ".join(r.cognitive_capability), r.priority, r.preliminary_asil, r.asil_note,
         ", ".join(r.standard_references), r.verification_method, r.acceptance_criteria, r.change_note,
         f"{r.updated_at} {r.updated_by}".strip()] for r in h.requirements], first=True)
    _sheet(wb, "Traceability", ["Requirement ID", "Requirement title", "Finding ID", "Agent", "Finding title", "Source"],
           [[t["requirement_id"], t["requirement_title"], t["finding_id"], t["agent"], t["finding_title"], t["source"]]
            for t in h.traceability])
    ws = _sheet(wb, "Cognitive coverage", ["Capability ID", "Faculty", "Capability", "Level", "Covered", "Requirement IDs"],
                [[c["id"], c["faculty"], c["title"], c["level"], "yes" if c["covered"] else "no",
                  ", ".join(c["requirement_ids"])] for c in h.coverage])
    ws.insert_rows(1)
    ws["A1"] = f"Target {h.target_level}; specified {h.specified_level}"
    ws["A1"].font = Font(bold=True)
    _sheet(wb, "Open items", ["Kind", "ID", "State", "Summary", "Owner hint"],
           [[o["kind"], o["id"], str(o["state"]), o["summary"], o["owner"]] for o in h.open_items])
    _sheet(wb, "Run log", ["Run ID", "Agent", "Mode", "Provider", "Model", "Turns", "Input tokens", "Output tokens",
                           "Status", "Started", "Finished", "Summary"],
           [[r["run_id"], r["agent_id"], r["mode"], r["provider"], r["model"], r["turns"], r["input_tokens"],
             r["output_tokens"], r["status"], r["started"], r["finished"], r["summary"]] for r in h.run_log])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def filename(h: Handoff) -> str:
    return f"SEW1-IVA-requirements-{h.stamp()}.xlsx"
