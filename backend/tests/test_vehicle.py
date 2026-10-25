"""Vehicle context and GIVEN/WHEN/THEN scenarios. See docs/12-vehicle-context.md."""

from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import orchestrator as orch, register, vehicle
from app.knowledge import load_pack
from app.main import app
from app.models import Requirement, State
from app.orchestrator import Orchestrator

COLD = {
    "given": [{"signal": "ignition_state", "op": "==", "value": "ON"},
              {"signal": "cabin_temp", "op": "<", "value": "@setpoint_driver"}],
    "when": {"trigger": "UTTERANCE", "intent": "thermal_discomfort.cold", "example": "I'm freezing"},
    "then": [{"type": "ACT", "target": "climate_setpoint", "zone": "speaker", "value": "+2 °C", "say": "Warming up."}],
}


def checks(scenario):
    return {i["check"]: i["severity"] for i in vehicle.validate(scenario)}


def test_every_shipped_scenario_is_valid():
    packs = ["vehicle_scenarios", "companion_cognition", "production_assistants", "distraction", "safety_security"]
    for name in packs:
        data = load_pack(name)
        for e in data.get("entries") or data["capabilities"]:
            if sc := (e.get("proposal") or {}).get("scenario"):
                assert vehicle.validate(sc) == [], (name, e.get("key") or e.get("id"))


def test_validate_flags_unknown_signal_and_enum_value():
    bad = {**COLD, "given": [{"signal": "cabin_tmp", "op": "<", "value": 19},
                             {"signal": "ignition_state", "op": "==", "value": True}]}
    assert checks(bad) == {"unknown-signal": "major", "enum-value": "major"}


def test_high_impact_needs_confirmation_and_prohibited_needs_refusal():
    window = {**COLD, "then": [{"type": "ACT", "target": "window_pos_driver", "value": 0}]}
    adas = {**COLD, "then": [{"type": "CONFIRM_THEN_ACT", "target": "adas_cruise_active", "value": True}]}
    assert checks(window)["high-impact"] == "critical"
    assert checks(adas)["prohibited"] == "critical"


def test_proactive_speech_while_moving_needs_workload():
    moving = {"given": [{"signal": "ac_on", "op": "==", "value": False}],
              "when": {"trigger": "THERMAL_DEVIATION_SUSTAINED"},
              "then": [{"type": "SUGGEST", "target": "ac_on", "value": True, "say": "AC?"}]}
    assert checks(moving) == {"workload": "major"}
    with_workload = {**moving, "given": moving["given"] + [{"signal": "driver_workload", "op": "==", "value": "low"}]}
    parked = {**moving, "when": {"trigger": "GEAR_PARK"}}
    assert vehicle.validate(with_workload) == [] and vehicle.validate(parked) == []


def test_render_reads_as_given_when_then():
    text = vehicle.render(COLD)
    assert text.startswith("GIVEN ignition state is ON AND cabin temperature < driver climate setpoint")
    assert "WHEN the occupant says “I'm freezing”" in text
    assert "THEN ACT climate setpoint [speaker zone] → +2 °C: “Warming up.”" in text


def test_vehicle_context_agent_validates_before_submitting():
    from app.runtime.events import list_events
    Orchestrator().run_cycle()
    run = orch.latest_run("vehicle_context")
    names = [e["data"]["name"] for e in list_events(run["run_id"]) if e["type"] == "tool_call"]
    assert names.index("check_scenarios") < names.index("submit_result")
    props = run["result"]["proposals"]
    assert len(props) >= 20 and all(p["scenario"] and vehicle.validate(p["scenario"]) == [] for p in props)


def test_gate_flags_a_bad_scenario_as_critical():
    bad = {**COLD, "then": [{"type": "ACT", "target": "sunroof_pos", "value": 0}]}
    req = Requirement(id="SYS-COG-900", title="Bad", statement="s", rationale="r", source="f", category="SYS-COG",
                      cognitive_capability=["COG-02"], verification_method="Test", acceptance_criteria="a",
                      scenario=bad, state=State.draft)
    register._insert(req)
    o = Orchestrator()
    o.run_cycle()
    gate = orch.latest_result("compliance_gate")
    assert any(f.title == "SYS-COG-900 scenario: high-impact" and f.confidence == "critical" for f in gate.findings)


def test_full_cycle_with_scenarios_reaches_cl5_cleanly_and_exports():
    o = Orchestrator()
    o.run_cycle()
    for p in register.proposals("open"):
        register.accept(p["pk"])
    o.run_cycle()
    gate = orch.latest_result("compliance_gate")
    assert gate.metrics["specified_cognitive_level"] == "CL5" and gate.findings == []
    with TestClient(app) as client:
        sig = client.get("/api/signals").json()
        assert len(sig["signals"]) >= 70 and all(t["used_by"] for t in sig["triggers"])
        wb = load_workbook(BytesIO(client.get("/api/export/excel").content))
        assert wb["Scenarios"].max_row > 30
        doc = Document(BytesIO(client.get("/api/export/word").content))
        assert any(p.text.startswith("GIVEN ") for p in doc.paragraphs)
