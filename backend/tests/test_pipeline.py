"""W0 exit criteria: docs/10-plan.md#workbench-track."""

from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import orchestrator as orch, register
from app.main import app
from app.models import RunStatus, State
from app.orchestrator import DependencyError, Orchestrator


def accept_all(state=State.draft):
    for p in register.proposals("open"):
        register.accept(p["pk"], state)


def test_full_cycle_produces_inbox_and_gate_level():
    o = Orchestrator()
    run_ids = o.run_cycle()
    assert len(run_ids) == 9
    assert all(r["status"] == RunStatus.succeeded for r in orch.runs())
    inbox = register.proposals("open")
    assert inbox and {p["agent_id"] for p in inbox} <= {
        "requirements_synthesis", "gdpr_privacy", "driver_distraction", "safety_cybersecurity"}
    gate = orch.latest_result("compliance_gate")
    assert gate.metrics["specified_cognitive_level"] == "CL0"
    assert len(gate.metrics["uncovered"]) == 20


def test_accepting_proposals_reaches_target_cl5():
    o = Orchestrator()
    o.run_cycle()
    accept_all()
    o.run_cycle()
    gate = orch.latest_result("compliance_gate")
    assert gate.metrics["specified_cognitive_level"] == "CL5"
    assert gate.metrics["uncovered"] == []


def test_rerun_does_not_duplicate_register_items():
    o = Orchestrator()
    o.run_cycle()
    accept_all()
    o.run_cycle()
    titles = [p["proposal"]["title"] for p in register.proposals("open")]
    assert titles == [], titles


def test_synthesis_refuses_without_research():
    with pytest.raises(DependencyError):
        Orchestrator().start("requirements_synthesis")


def test_safety_proposals_keep_asil_note():
    Orchestrator().run_cycle()
    safe = [p["proposal"] for p in register.proposals("open") if p["proposal"]["category"] == "SYS-SAFE"]
    assert safe and all(p["asil_note"] == "to be confirmed by the safety manager" for p in safe)


def test_accept_new_version_supersedes_and_keeps_id():
    req = register.add({"category": "SYS-COG", "title": "Trip memory", "statement": "Old.",
                        "cognitive_capability": ["COG-07"]})
    register.store_proposals("run_x", "requirements_synthesis", [
        register.Proposal(id=req.id, title="Trip memory", statement="New.", rationale="r", source="find_cog_07",
                          category="SYS-COG")])
    pk = register.proposals("open")[0]["pk"]
    new = register.accept(pk, State.in_review)
    assert new.id == req.id and new.version == 2
    assert register.history(req.id)[0].state == State.superseded


def test_only_baseline_roles_may_baseline():
    register.store_proposals("run_x", "gdpr_privacy", [
        register.Proposal(title="t", statement="s", rationale="r", source="f", category="SYS-PRIV")])
    pk = register.proposals("open")[0]["pk"]
    with pytest.raises(register.RegisterError):
        register.accept(pk, State.baselined, role="systems_engineer")
    assert register.accept(pk, State.baselined, role="baseline_owner").state == State.baselined


def test_api_and_exports():
    o = Orchestrator()
    o.run_cycle()
    accept_all()
    with TestClient(app) as client:
        assert len(client.get("/api/agents").json()) == 9
        assert client.get("/api/cognition").json()["specified"] == "CL5"
        xlsx = client.get("/api/export/excel")
        assert xlsx.status_code == 200
        wb = load_workbook(BytesIO(xlsx.content))
        assert wb.sheetnames == ["Requirements", "Traceability", "Cognitive coverage", "Open items", "Run log"]
        docx = client.get("/api/export/word")
        text = "\n".join(p.text for p in Document(BytesIO(docx.content)).paragraphs)
        assert "6. Cognitive companion" in text and "target CL5, specified CL5" in text


def test_agents_run_through_the_model_loop_with_a_trace():
    from app.runtime.events import list_events
    Orchestrator().run_cycle()
    run = orch.latest_run("requirements_synthesis", RunStatus.succeeded)
    types = [e["type"] for e in list_events(run["run_id"])]
    assert types[0] == "start" and types[-1] == "result"
    assert types.count("tool_call") >= 7  # four upstream fetches, register, inbox, submit
    metrics = run["result"]["metrics"]
    assert metrics["provider"] == "simulated" and metrics["turns"] == 3 and metrics["input_tokens"] > 0


def test_simulated_policies_only_see_tool_results():
    """A policy given an empty transcript must ask for tools, not answer from hidden state."""
    from app.agents import AGENTS
    from app.runtime.simulation import POLICIES
    from app.runtime.tools import TOOLS
    from app.runtime.types import Message, ModelRequest
    for agent in AGENTS.values():
        req = ModelRequest(agent_id=agent.id, model="m", system="s", messages=[Message(role="user", text="go")],
                           tools=[TOOLS[t].spec for t in agent.tools])
        calls = POLICIES[agent.id](req).tool_calls
        assert calls and all(c.name in agent.tools and c.name != "submit_result" for c in calls)


def test_tool_errors_are_returned_to_the_model():
    from app.models import RunContext
    from app.runtime.tools import get_dependency_output
    ctx = RunContext(run_id="r", agent_id="a", project={}, register=[], dependency_runs=[])
    with pytest.raises(ValueError):
        get_dependency_output(ctx, "requirements_synthesis")


def test_unselected_provider_fails_the_run_cleanly(monkeypatch):
    from app import config
    monkeypatch.setattr(config, "runtime", lambda: {"provider": "anthropic", "model": "x"})
    o = Orchestrator()
    o.start("google_assistant_reference")
    wait_all(o)
    run = orch.latest_run("google_assistant_reference")
    assert run["status"] == RunStatus.failed and "not implemented" in run["summary"]


def wait_all(o):
    from concurrent.futures import wait
    wait(list(o.futures.values()))
